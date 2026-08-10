#!/usr/bin/env python3
"""Приём .xlsx от носителя языка обратно в словарь.

    python3 tools/xlsx_import.py TPI_LUX_перевод_KA.xlsx [--dry-run]

Читает колонку «ქართული» с листа «Перевод» и колонку ka с листа «Термины»,
кладёт их в src/i18n/strings.json и src/i18n/glossary.json.
Комментарии переводчика складывает в docs/native-review-notes.md.

Дальше: python3 tools/check.py && python3 tools/build.py
"""
import json
import sys
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parent.parent
STRINGS = ROOT / "src" / "i18n" / "strings.json"
GLOSSARY = ROOT / "src" / "i18n" / "glossary.json"
NOTES = ROOT / "docs" / "native-review-notes.md"

# колонки ищем по началу заголовка — так файл переживает перестановку колонок
HEADERS = {
    "id": "ID",
    "group": "Раздел сайта",
    "where": "Где на странице",
    "src": "Русский",
    "ours": "ქართული — наш перевод",
    "fix": "✍ ПРАВКА",
    "note": "Комментарий",
}


def locate(ws):
    found = {}
    for col in range(1, ws.max_column + 1):
        title = str(ws.cell(row=1, column=col).value or "").strip()
        for name, prefix in HEADERS.items():
            if name not in found and title.startswith(prefix):
                found[name] = col
    missing = [n for n in ("id", "ours", "fix") if n not in found]
    if missing:
        sys.exit(f"на листе «Перевод» не нашлись колонки: {', '.join(missing)}")
    return found


def cell(ws, row, col):
    if not col:
        return ""
    value = ws.cell(row=row, column=col).value
    return str(value).strip() if value is not None else ""


def import_strings(ws, data, notes):
    col = locate(ws)
    accepted = corrected = same = 0
    unknown, empty = [], []
    for row in range(2, ws.max_row + 1):
        key = cell(ws, row, col["id"])
        if not key:
            continue
        if key not in data:
            unknown.append((row, key))
            continue
        ours = cell(ws, row, col["ours"])
        fix = cell(ws, row, col["fix"])
        note = cell(ws, row, col.get("note"))
        ka = fix or ours          # правка носителя всегда важнее нашего перевода
        if note:
            notes.append((cell(ws, row, col.get("group")), cell(ws, row, col.get("src")), ka, note))
        if not ka:
            empty.append(key)
            continue
        old = (data[key].get("ka") or "").strip()
        if fix and fix != ours:
            corrected += 1
        elif old == ka:
            same += 1
        else:
            accepted += 1
        data[key]["ka"] = ka
        # правка носителя закрывает строку; принятый машинный перевод — ещё нет
        data[key]["st_ka"] = "native" if fix else "mt-approved"
    return {
        "accepted": accepted,
        "corrected": corrected,
        "same": same,
        "empty": empty,
        "unknown": unknown,
    }


def import_glossary(ws, glossary, notes):
    terms = {t.get("ru", ""): t for t in glossary.get("terms", [])}
    changed = 0
    for row in range(2, ws.max_row + 1):
        ru = cell(ws, row, 1)
        ka = cell(ws, row, 2)
        note = cell(ws, row, 5)
        term = terms.get(ru)
        if not term or not ka:
            continue
        if note and note not in (
            "перевод устоявшийся — беглая проверка",
            "нужно решение носителя",
            "уже вычитано",
        ):
            notes.append(("Термины", ru, ka, note))
        if (term.get("ka") or "").strip() != ka:
            changed += 1
        term["ka"] = ka
        term["status"] = "done"
    return changed


def write_notes(notes):
    if not notes:
        return
    lines = ["# Замечания носителя языка", ""]
    group = None
    for section, src, ka, note in notes:
        if section != group:
            lines += [f"## {section}", ""]
            group = section
        lines += [f"- **{src}**", f"  - ka: {ka or '—'}", f"  - замечание: {note}", ""]
    NOTES.write_text("\n".join(lines), encoding="utf-8")


def main():
    args = [a for a in sys.argv[1:] if not a.startswith("--")]
    dry = "--dry-run" in sys.argv
    if not args:
        sys.exit("укажите путь к .xlsx")
    path = Path(args[0])
    if not path.exists():
        sys.exit(f"нет файла: {path}")

    wb = load_workbook(path, data_only=True)
    if "Перевод" not in wb.sheetnames:
        sys.exit(f"в книге нет листа «Перевод» (есть: {', '.join(wb.sheetnames)})")

    data = json.loads(STRINGS.read_text(encoding="utf-8"))
    glossary = json.loads(GLOSSARY.read_text(encoding="utf-8"))
    notes = []

    stats = import_strings(wb["Перевод"], data, notes)
    terms_changed = import_glossary(wb["Термины"], glossary, notes) if "Термины" in wb.sheetnames else 0

    print(f"исправлено носителем:  {stats['corrected']}")
    print(f"наш перевод принят:    {stats['accepted'] + stats['same']}")
    print(f"осталось пустых:       {len(stats['empty'])}")
    print(f"терминов обновлено:    {terms_changed}")
    print(f"замечаний:             {len(notes)}")
    if stats["unknown"]:
        print(f"\n⚠ строк с неизвестным ID (пропущены): {len(stats['unknown'])}")
        for row, key in stats["unknown"][:10]:
            print(f"   строка {row}: {key}")

    if dry:
        print("\n--dry-run: файлы не тронуты")
        return

    STRINGS.write_text(json.dumps(data, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    GLOSSARY.write_text(json.dumps(glossary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    write_notes(notes)
    print(f"\nзаписано: {STRINGS.relative_to(ROOT)}, {GLOSSARY.relative_to(ROOT)}"
          + (f", {NOTES.relative_to(ROOT)}" if notes else ""))
    print("дальше:   python3 tools/check.py && python3 tools/build.py")


if __name__ == "__main__":
    main()
